"""Approved adapters retain parser identity and source/table locations."""

from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from rag_service.services.parsing import parse, table_segments


@pytest.mark.parametrize(
    ("filename", "content", "parser", "location"),
    [
        ("readme.md", b"# Overview\nUseful text", "text", "section"),
        ("notes.txt", b"plain text", "text", "line_start"),
        ("src/main.py", b"def main():\n    return 1", "text", "source_path"),
        ("data.csv", b"name,value\nalpha,1\nbeta,2\n,,", "csv", "row_start"),
        (
            "page.html",
            b"<h1>Title</h1><p>Text</p><script>secret()</script>",
            "beautifulsoup4",
            "section",
        ),
        ("page.htm", b"standalone", "beautifulsoup4", None),
    ],
)
def test_text_and_html_adapters(
    filename: str, content: bytes, parser: str, location: str | None
) -> None:
    name, revision, parts = parse(content, filename)
    assert name == parser
    assert revision
    assert parts
    assert "secret()" not in str(parts)
    if location:
        assert location in parts[-1]["provenance"]  # type: ignore[operator]


def test_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_heading("Section", level=1)
    document.add_paragraph("Some words")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Alpha"
    table.cell(1, 1).text = "1"
    buffer = BytesIO()
    document.save(buffer)
    name, _, parts = parse(buffer.getvalue(), "report.docx")
    assert name == "python-docx"
    assert parts[1]["provenance"] == {"section": "Section", "paragraph": 2}
    assert parts[-1]["provenance"]["columns"] == ["Name", "Value"]  # type: ignore[index]


def test_xlsx_sheets_headers_and_blank_cells() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Results"
    sheet.append(["Name", "Value"])
    sheet.append(["Alpha", 2])
    sheet.append(["Beta", None])
    empty = workbook.create_sheet("Empty")
    empty.append(["Header"])
    buffer = BytesIO()
    workbook.save(buffer)
    name, _, parts = parse(buffer.getvalue(), "report.xlsx")
    assert name == "openpyxl"
    assert parts[0]["provenance"]["sheet"] == "Results"  # type: ignore[index]
    assert parts[1]["text"] == "Beta | "
    assert table_segments([], "Empty") == []


def test_pdf_page_location_and_binary_failures() -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 50 Td (Fixture text) Tj ET")
    page[NameObject("/Contents")] = stream
    buffer = BytesIO()
    writer.write(buffer)
    name, _, parts = parse(buffer.getvalue(), "report.pdf")
    assert name == "pypdf"
    assert parts[0]["provenance"] == {"page": 1}
    assert "Fixture text" in str(parts[0]["text"])
    with pytest.raises(ValueError, match="Unsupported binary"):
        parse(b"abc\x00def", "binary.bin")
    with pytest.raises(UnicodeDecodeError):
        parse(b"\xff", "text.txt")
