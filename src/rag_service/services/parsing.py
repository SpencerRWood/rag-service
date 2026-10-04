"""Approved format adapters produce normalized segments without model inference."""

import csv
import io
from importlib.metadata import version
from pathlib import PurePosixPath

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def segment(text: str, **provenance: object) -> dict[str, object]:
    """One source unit with independently preserved location/context."""
    return {"text": text, "provenance": provenance}


def table_segments(rows: list[list[str]], sheet: str) -> list[dict[str, object]]:
    """Keep headers and physical row/column locations on every data row."""
    if not rows:
        return []
    headers = rows[0]
    return [
        segment(
            " | ".join(row),
            sheet=sheet,
            row_start=number,
            row_end=number,
            columns=headers,
            column_start=1,
            column_end=len(row),
            table_header=" | ".join(headers),
        )
        for number, row in enumerate(rows[1:], start=2)
        if any(row)
    ] or [segment(" | ".join(headers), sheet=sheet, row_start=1, columns=headers)]


def parse(  # noqa: PLR0912 -- explicit approved format dispatch
    content: bytes, filename: str
) -> tuple[str, str, list[dict[str, object]]]:
    """Return parser identity/version and normalized source units.

    Unsupported binary formats fail explicitly. Text/source-code extensions use
    strict UTF-8 (including BOM); no lossy replacement of source bytes occurs.
    """
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content), strict=True)
        return (
            "pypdf",
            version("pypdf"),
            [
                segment(page.extract_text() or "", page=number)
                for number, page in enumerate(reader.pages, start=1)
            ],
        )
    if suffix == ".docx":
        document = Document(io.BytesIO(content))
        parts: list[dict[str, object]] = []
        section = ""
        for number, paragraph in enumerate(document.paragraphs, start=1):
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                section = paragraph.text
            parts.append(segment(paragraph.text, section=section, paragraph=number))
        for number, table in enumerate(document.tables, start=1):
            parts.extend(
                table_segments(
                    [[cell.text for cell in row.cells] for row in table.rows],
                    f"table-{number}",
                )
            )
        return "python-docx", version("python-docx"), parts
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            parts = []
            for sheet in workbook:
                rows = [
                    ["" if value is None else str(value) for value in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                parts.extend(table_segments(rows, sheet.title))
            return "openpyxl", version("openpyxl"), parts
        finally:
            workbook.close()
    text = content.decode("utf-8-sig")
    if "\x00" in text:
        raise ValueError("Unsupported binary source")
    if suffix == ".csv":
        rows = list(csv.reader(io.StringIO(text)))
        return "csv", "stdlib-1", table_segments(rows, filename)
    if suffix in {".html", ".htm"}:
        soup = BeautifulSoup(text, "html.parser")
        for element in soup(["script", "style"]):
            element.decompose()
        parts = []
        section = ""
        for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "tr"]):
            value = element.get_text(" ", strip=True)
            if element.name.startswith("h"):
                section = value
            parts.append(segment(value, section=section))
        return (
            "beautifulsoup4",
            version("beautifulsoup4"),
            parts or [segment(soup.get_text(" ", strip=True))],
        )
    parts = []
    section = ""
    lines: list[str] = []
    start = 1
    for number, line in enumerate(text.splitlines(), start=1):
        if suffix in {".md", ".markdown"} and line.startswith("#"):
            if lines:
                parts.append(
                    segment(
                        "\n".join(lines),
                        section=section,
                        source_path=filename,
                        line_start=start,
                        line_end=number - 1,
                    )
                )
            lines = []
            start = number
            section = line.lstrip("# ")
        lines.append(line)
    if lines:
        parts.append(
            segment(
                "\n".join(lines),
                section=section,
                source_path=filename,
                line_start=start,
                line_end=len(text.splitlines()),
            )
        )
    return "text", "utf8-structural-1", parts
